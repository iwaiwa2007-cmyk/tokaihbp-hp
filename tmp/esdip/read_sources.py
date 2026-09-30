from pathlib import Path
from pypdf import PdfReader
from docx import Document
import openpyxl
root=Path('/Users/iwasakieiten/Desktop/20260929　十二指腸ESDIPの申請書類')
out=Path('tmp/esdip')
for f in root.glob('*.pdf'):
 r=PdfReader(f); s='\n'.join(f'\nPAGE {i+1}\n'+(p.extract_text() or '') for i,p in enumerate(r.pages)); (out/(f.stem+'.txt')).write_text(s)
 print(f.name,len(r.pages),len(s))
for f in (root/'高難度新規医療技術の申請書類、手続き').glob('*'):
 if f.suffix=='.docx':
  d=Document(f); s='\n'.join(p.text for p in d.paragraphs)
  for i,t in enumerate(d.tables):
   s+=f'\nTABLE {i}\n'+'\n'.join(' | '.join(c.text for c in row.cells) for row in t.rows)
  (out/(f.stem+'.txt')).write_text(s)
  print(f.name,s)
 elif f.suffix=='.xlsx':
  w=openpyxl.load_workbook(f); s=''
  for sh in w:
   s+='\nSHEET '+sh.title+'\n'
   for row in sh:
    s+=' '.join(f'{c.coordinate}={c.value}' for c in row if c.value is not None)+'\n'
  (out/(f.stem+'.txt')).write_text(s)
  print(f.name,s)
